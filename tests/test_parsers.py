from __future__ import annotations

from pathlib import Path

import fitz
from docx import Document as DocxDocument

from app.parsers import parse_file


def test_parse_multiple_formats(tmp_path: Path) -> None:
    md_path = tmp_path / "guide.md"
    md_path.write_text("# Title\n\n## Section\nCheck upstream health.", encoding="utf-8")

    txt_path = tmp_path / "guide.txt"
    txt_path.write_text("First paragraph.\n\nSecond paragraph.", encoding="utf-8")

    html_path = tmp_path / "guide.html"
    html_path.write_text("<html><head><title>HTML Guide</title></head><body><h1>HTML Guide</h1><p>Check logs.</p></body></html>", encoding="utf-8")

    docx_path = tmp_path / "guide.docx"
    doc = DocxDocument()
    doc.add_heading("DOCX Guide", level=1)
    doc.add_paragraph("Validate configuration.")
    doc.save(str(docx_path))

    pdf_path = tmp_path / "guide.pdf"
    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "PDF Guide\nInspect service logs.")
    pdf.save(pdf_path)
    pdf.close()

    for path in (md_path, txt_path, html_path, docx_path, pdf_path):
        title, source_type, blocks = parse_file(path)
        assert title
        assert source_type
        assert blocks
        assert any(block.content for block in blocks)
